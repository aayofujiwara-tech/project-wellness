#!/usr/bin/env python3
"""PDF利用者情報 → Excel ヒアリングシート 自動転記スクリプト。

input_pdf/ 内のPDFファイルを解析し、利用者氏名を基に excel_sheets/ 内の
Excelヒアリングシートへデータを自動転記する。

運用方針:
    デフォルトで上段フォーム（form1）に書き込む。
    下段（form2）は --form form2 で切り替え可能。

使い方:
    python auto_pdf_to_excel.py                          # 上段(form1)に転記
    python auto_pdf_to_excel.py --dry-run                # 書き込みせずに確認のみ
    python auto_pdf_to_excel.py --form form2             # 下段(form2)に転記
    python auto_pdf_to_excel.py --config config.json --pdf-dir input_pdf --excel-dir excel_sheets
"""

import argparse
import json
import logging
import os
import re
import shutil
import sys
import unicodedata
from datetime import datetime

import fitz  # pymupdf
from openpyxl import load_workbook

# ---------------------------------------------------------------------------
# ロギング設定
# ---------------------------------------------------------------------------
LOG_FORMAT = "%(asctime)s [%(levelname)s] %(message)s"
DATE_FORMAT = "%Y-%m-%d %H:%M:%S"


def setup_logging(log_file=None):
    """ロギングを設定する。"""
    handlers = [logging.StreamHandler(sys.stdout)]
    if log_file:
        handlers.append(logging.FileHandler(log_file, encoding="utf-8"))
    logging.basicConfig(level=logging.INFO, format=LOG_FORMAT, datefmt=DATE_FORMAT, handlers=handlers)


# ---------------------------------------------------------------------------
# 設定読み込み
# ---------------------------------------------------------------------------
def load_config(config_path):
    """config.json を読み込む。"""
    with open(config_path, "r", encoding="utf-8") as f:
        return json.load(f)


# ---------------------------------------------------------------------------
# テキスト正規化
# ---------------------------------------------------------------------------
def normalize_text(text):
    """全角英数字・記号を半角に正規化する。"""
    return unicodedata.normalize("NFKC", text)


# ---------------------------------------------------------------------------
# バックアップ・ファイル整理
# ---------------------------------------------------------------------------
def backup_excel(filepath, backup_base_dir):
    """Excelファイルをバックアップフォルダへコピーする。"""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_dir = os.path.join(backup_base_dir, timestamp)
    os.makedirs(backup_dir, exist_ok=True)
    dest = os.path.join(backup_dir, os.path.basename(filepath))
    shutil.copy2(filepath, dest)
    return dest


def move_processed_pdf(pdf_path, processed_base_dir):
    """処理済みPDFを processed/YYYYMMDD/ フォルダへ移動する。"""
    date_str = datetime.now().strftime("%Y%m%d")
    dest_dir = os.path.join(processed_base_dir, date_str)
    os.makedirs(dest_dir, exist_ok=True)
    dest = os.path.join(dest_dir, os.path.basename(pdf_path))
    if os.path.exists(dest):
        base, ext = os.path.splitext(os.path.basename(pdf_path))
        counter = 1
        while os.path.exists(dest):
            dest = os.path.join(dest_dir, f"{base}_{counter}{ext}")
            counter += 1
    shutil.move(pdf_path, dest)
    return dest


# ---------------------------------------------------------------------------
# PDF 解析
# ---------------------------------------------------------------------------
def extract_text_from_pdf(pdf_path):
    """PDFからテキストを抽出する（pymupdf使用）。"""
    all_text = []
    doc = fitz.open(pdf_path)
    for page in doc:
        text = page.get_text()
        if text:
            all_text.append(text)
    doc.close()
    return "\n".join(all_text)


def parse_pdf_text(text, patterns):
    """抽出テキストから正規表現でフィールドを抽出し、辞書に構造化する。"""
    normalized = normalize_text(text)
    result = {}
    for field_name, regex_list in patterns.items():
        for pattern in regex_list:
            normalized_pattern = normalize_text(pattern)
            match = re.search(normalized_pattern, normalized)
            if match:
                value = match.group(1).strip()
                result[field_name] = value
                break
    return result


def extract_data_from_pdf(pdf_path, config):
    """PDFファイル1つから構造化データを抽出する。"""
    text = extract_text_from_pdf(pdf_path)
    patterns = config["pdf_extraction_patterns"]
    return parse_pdf_text(text, patterns)


def get_surname_from_pdf_data(data):
    """抽出データから苗字を取得する。"""
    full_name = data.get("利用者氏名", "")
    if not full_name:
        return None
    parts = re.split(r"[\s　]+", full_name)
    return parts[0] if parts else None


# ---------------------------------------------------------------------------
# Excel ファイル走査・照合
# ---------------------------------------------------------------------------
def find_excel_files(excel_dir):
    """ディレクトリ内のヒアリングシートExcelファイルを一覧する。"""
    files = []
    for f in os.listdir(excel_dir):
        if f.endswith(".xlsx") and "ヒアリングシート" in f:
            files.append(os.path.join(excel_dir, f))
    return sorted(files)


def extract_name_from_filename(filename, config):
    """ファイル名から氏名（苗字）部分を抽出する。"""
    pattern = config.get("filename_pattern", "")
    match = re.search(pattern, os.path.basename(filename))
    if match:
        raw = match.group(1)
        abbrevs = config.get("facility_abbreviations", {})
        for abbr in abbrevs:
            if raw.startswith(abbr):
                return raw[len(abbr):]
        return raw
    return None


def match_excel_file(surname, excel_files, config):
    """苗字に一致するExcelファイルを部分一致で探す。"""
    matches = []
    for filepath in excel_files:
        file_name_part = extract_name_from_filename(filepath, config)
        if file_name_part and surname in file_name_part:
            matches.append(filepath)
    return matches


# ---------------------------------------------------------------------------
# 要介護度の正規化
# ---------------------------------------------------------------------------
def normalize_care_level(value, config):
    """PDF抽出値をconfig内の正規名称に正規化する。"""
    normalized_value = normalize_text(value)
    care_map = config.get("care_level_mapping", {})
    for canonical, aliases in care_map.items():
        normalized_aliases = [normalize_text(a) for a in aliases]
        if normalized_value in normalized_aliases or normalized_value == canonical:
            return canonical
    return normalized_value


# ---------------------------------------------------------------------------
# フォーマット関数群
# ---------------------------------------------------------------------------
def format_gender(value):
    """性別を短縮形に変換する。男性→男, 女性→女"""
    normalized = normalize_text(value.strip())
    if "男" in normalized:
        return "男"
    elif "女" in normalized:
        return "女"
    return normalized


def format_age(value):
    """年齢に「歳」を付与する。既に「歳」があればそのまま。"""
    normalized = normalize_text(str(value).strip())
    age_match = re.search(r"(\d+)", normalized)
    if age_match:
        return f"{age_match.group(1)}歳"
    return normalized


def format_furigana(value):
    """ふりがなの括弧を除去する。"""
    cleaned = re.sub(r"[（()）]", "", value)
    return cleaned.strip()


def format_care_level_inline(value, config):
    """要介護度を丸数字に変換してインライン文字列を返す。

    例: "要介護3" → "要支援（ １・２ ）　／　要介護（ 1・2・③・4・5 ）"
    """
    canonical = normalize_care_level(value, config)
    circled = {"1": "①", "2": "②", "3": "③", "4": "④", "5": "⑤"}

    support_match = re.match(r"要支援(\d)", canonical)
    care_match = re.match(r"要介護(\d)", canonical)

    if support_match:
        num = support_match.group(1)
        support_part = "１・２"
        if num == "1":
            support_part = "①・２"
        elif num == "2":
            support_part = "１・②"
        return f"要支援（ {support_part} ）　／　要介護（ 1・2・3・4・5 ）"

    elif care_match:
        num = care_match.group(1)
        nums_list = ["1", "2", "3", "4", "5"]
        formatted_nums = []
        for n in nums_list:
            if n == num:
                formatted_nums.append(circled[n])
            else:
                formatted_nums.append(n)
        care_part = "・".join(formatted_nums)
        return f"要支援（ １・２ ）　／　要介護（ {care_part} ）"

    return canonical


# ---------------------------------------------------------------------------
# Excel 座標指定書き込み (fixed_cell_mapping 方式)
# ---------------------------------------------------------------------------
def write_data_to_excel(filepath, data, config, form_name="form1", dry_run=False):
    """固定座標指定方式でExcelにデータを書き込む。

    config["fixed_cell_mapping"][form_name] のセル座標に直接書き込む。
    空値(None)の場合は既存データを保護し上書きしない。
    """
    cell_mapping = config["fixed_cell_mapping"][form_name]
    target_sheet = config.get("target_sheet")
    wb = load_workbook(filepath)

    if target_sheet and target_sheet in wb.sheetnames:
        ws = wb[target_sheet]
    else:
        ws = wb.active

    format_functions = {
        "format_gender": format_gender,
        "format_age": format_age,
        "format_furigana": format_furigana,
        "format_care_level_inline": lambda v: format_care_level_inline(v, config),
    }

    written_fields = []
    skipped_fields = []
    protected_fields = []

    for field_name, field_info in cell_mapping.items():
        value = data.get(field_name)
        cell_addr = field_info["cell"]

        # 空欄保護: 抽出結果がNoneまたは空文字の場合は既存値を保持
        if value is None or (isinstance(value, str) and value.strip() == ""):
            existing = ws[cell_addr].value
            if existing and str(existing).strip():
                protected_fields.append((field_name, str(existing)))
            else:
                skipped_fields.append(field_name)
            continue

        # フォーマット関数を適用
        fmt = field_info.get("format")
        if fmt and fmt in format_functions:
            value = format_functions[fmt](value)

        if not dry_run:
            ws[cell_addr].value = value

        written_fields.append((field_name, value, cell_addr))

    if not dry_run:
        wb.save(filepath)

    wb.close()
    return written_fields, skipped_fields, protected_fields


# ---------------------------------------------------------------------------
# レポート出力
# ---------------------------------------------------------------------------
def generate_report(report_path, success_list, no_name_pdfs, no_match_pdfs,
                    multi_match_pdfs, unmatched_excels, backup_log, move_log,
                    protected_log, dry_run=False):
    """簡易レポートファイルを生成する。"""
    mode_label = "【ドライラン】" if dry_run else ""
    lines = []
    lines.append(f"{mode_label}PDF → Excel 自動転記 実行レポート")
    lines.append(f"生成日時: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    lines.append("=" * 60)

    lines.append(f"\n■ 成功: {len(success_list)} 件")
    for pdf_name, excel_name, count in success_list:
        lines.append(f"  {pdf_name} → {excel_name} ({count}項目転記)")

    if protected_log:
        lines.append(f"\n■ 空欄保護（既存値を維持）: {len(protected_log)} 件")
        for pdf_name, field, existing in protected_log:
            lines.append(f"  {pdf_name}: {field} (既存値「{existing}」を保護)")

    if no_name_pdfs:
        lines.append(f"\n■ 失敗（氏名抽出不可）: {len(no_name_pdfs)} 件")
        for name in no_name_pdfs:
            lines.append(f"  - {name}")

    if no_match_pdfs:
        lines.append(f"\n■ 失敗（対応Excel未発見）: {len(no_match_pdfs)} 件")
        for pdf_name, surname in no_match_pdfs:
            lines.append(f"  - {pdf_name} (苗字: {surname})")

    if multi_match_pdfs:
        lines.append(f"\n■ 注意（複数Excel一致）: {len(multi_match_pdfs)} 件")
        for pdf_name, excel_names in multi_match_pdfs:
            lines.append(f"  - {pdf_name} → {', '.join(excel_names)}")

    if unmatched_excels:
        lines.append(f"\n■ スキップ（対応PDFなしのExcel）: {len(unmatched_excels)} 件")
        for path in sorted(unmatched_excels):
            lines.append(f"  - {os.path.basename(path)}")

    if backup_log:
        lines.append(f"\n■ バックアップ: {len(backup_log)} 件")
        for src, dest in backup_log:
            lines.append(f"  {os.path.basename(src)} → {dest}")

    if move_log:
        lines.append(f"\n■ 処理済みPDF移動: {len(move_log)} 件")
        for src, dest in move_log:
            lines.append(f"  {os.path.basename(src)} → {dest}")

    lines.append("\n" + "=" * 60)

    content = "\n".join(lines)
    with open(report_path, "w", encoding="utf-8") as f:
        f.write(content)
    return report_path


# ---------------------------------------------------------------------------
# メイン処理
# ---------------------------------------------------------------------------
def process_all(config_path, pdf_dir, excel_dir, form_name="form1",
                dry_run=False, log_file=None):
    """全PDFを処理してExcelに転記する。"""
    setup_logging(log_file)
    logger = logging.getLogger(__name__)

    config = load_config(config_path)

    # フォーム名の存在チェック
    if form_name not in config.get("fixed_cell_mapping", {}):
        logger.error(f"フォーム「{form_name}」が config.json の fixed_cell_mapping に存在しません。")
        sys.exit(1)

    mode_label = "【ドライラン】" if dry_run else ""
    now = datetime.now()

    logger.info(f"{'='*60}")
    logger.info(f"{mode_label}PDF → Excel 自動転記処理 開始")
    logger.info(f"{'='*60}")
    logger.info(f"PDF入力元:   {os.path.abspath(pdf_dir)}")
    logger.info(f"Excel出力先: {os.path.abspath(excel_dir)}")
    logger.info(f"フォーム:    {form_name}")

    os.makedirs(pdf_dir, exist_ok=True)
    os.makedirs(excel_dir, exist_ok=True)

    pdf_files = [
        os.path.join(pdf_dir, f)
        for f in os.listdir(pdf_dir)
        if f.lower().endswith(".pdf")
    ]
    if not pdf_files:
        logger.warning("処理対象のPDFファイルが見つかりません。")
        return

    excel_files = find_excel_files(excel_dir)
    if not excel_files:
        logger.warning("対象のExcelファイルが見つかりません。")
        return

    logger.info(f"PDF数: {len(pdf_files)}, Excel数: {len(excel_files)}")

    backup_base = os.path.join(excel_dir, "backups")
    processed_base = os.path.join(excel_dir, "processed")

    success_list = []
    no_match_pdfs = []
    no_name_pdfs = []
    multi_match_pdfs = []
    unmatched_excels = set(excel_files)
    backup_log = []
    move_log = []
    protected_log = []
    processed_pdf_paths = []

    for pdf_path in sorted(pdf_files):
        pdf_name = os.path.basename(pdf_path)
        logger.info(f"\n--- 処理中: {pdf_name} ---")

        try:
            data = extract_data_from_pdf(pdf_path, config)
        except Exception as e:
            logger.error(f"  PDF解析エラー: {pdf_name} - {e}")
            no_name_pdfs.append(pdf_name)
            continue

        surname = get_surname_from_pdf_data(data)

        if not surname:
            logger.warning(f"  氏名を抽出できませんでした: {pdf_name}")
            no_name_pdfs.append(pdf_name)
            continue

        logger.info(f"  抽出氏名: {data.get('利用者氏名', '?')} (苗字: {surname})")

        matched = match_excel_file(surname, excel_files, config)

        if not matched:
            logger.warning(f"  一致するExcelファイルが見つかりません: 苗字「{surname}」")
            no_match_pdfs.append((pdf_name, surname))
            continue

        if len(matched) > 1:
            logger.warning(f"  複数のExcelが一致しました ({len(matched)}件)。全てに書き込みます。")
            multi_match_pdfs.append((pdf_name, [os.path.basename(f) for f in matched]))

        pdf_success = False
        for excel_path in matched:
            excel_name = os.path.basename(excel_path)

            if not dry_run:
                try:
                    backup_dest = backup_excel(excel_path, backup_base)
                    logger.info(f"  バックアップ: {excel_name} → {backup_dest}")
                    backup_log.append((excel_path, backup_dest))
                except Exception as e:
                    logger.error(f"  バックアップ失敗: {excel_name} - {e}")
                    continue

            logger.info(f"  転記先: {excel_name}")

            try:
                written, skipped, protected = write_data_to_excel(
                    excel_path, data, config,
                    form_name=form_name, dry_run=dry_run,
                )
            except Exception as e:
                logger.error(f"  Excel書き込みエラー: {excel_name} - {e}")
                continue

            for field, value, pos in written:
                logger.info(f"    ✓ {field}: {value} → {pos}")

            for field, existing in protected:
                logger.info(f"    🛡 {field}: 既存値「{existing}」を保護")
                protected_log.append((pdf_name, field, existing))

            if skipped:
                logger.info(f"    (スキップ: {', '.join(skipped)})")

            success_list.append((pdf_name, excel_name, len(written)))
            unmatched_excels.discard(excel_path)
            pdf_success = True

        if pdf_success:
            processed_pdf_paths.append(pdf_path)

    if not dry_run:
        for pdf_path in processed_pdf_paths:
            try:
                move_dest = move_processed_pdf(pdf_path, processed_base)
                logger.info(f"  PDF移動: {os.path.basename(pdf_path)} → {move_dest}")
                move_log.append((pdf_path, move_dest))
            except Exception as e:
                logger.error(f"  PDF移動失敗: {os.path.basename(pdf_path)} - {e}")

    logger.info(f"\n{'='*60}")
    logger.info(f"{mode_label}処理結果サマリー")
    logger.info(f"{'='*60}")

    logger.info(f"\n[成功] {len(success_list)} 件")
    for pdf_name, excel_name, count in success_list:
        logger.info(f"  {pdf_name} → {excel_name} ({count}項目転記)")

    if protected_log:
        logger.info(f"\n[空欄保護] {len(protected_log)} 件")
        for pdf_name, field, existing in protected_log:
            logger.info(f"  {pdf_name}: {field} (既存値「{existing}」を保護)")

    if no_name_pdfs:
        logger.warning(f"\n[未処理: 氏名抽出失敗] {len(no_name_pdfs)} 件")
        for name in no_name_pdfs:
            logger.warning(f"  - {name}")

    if no_match_pdfs:
        logger.warning(f"\n[未処理: Excel照合失敗] {len(no_match_pdfs)} 件")
        for pdf_name, surname in no_match_pdfs:
            logger.warning(f"  - {pdf_name} (苗字: {surname})")

    if multi_match_pdfs:
        logger.info(f"\n[注意: 複数Excel一致] {len(multi_match_pdfs)} 件")
        for pdf_name, excel_names in multi_match_pdfs:
            logger.info(f"  - {pdf_name} → {', '.join(excel_names)}")

    if unmatched_excels:
        logger.info(f"\n[情報: 対応PDFなしのExcel] {len(unmatched_excels)} 件")
        for path in sorted(unmatched_excels):
            logger.info(f"  - {os.path.basename(path)}")

    if backup_log:
        logger.info(f"\n[バックアップ] {len(backup_log)} 件作成済み")

    if move_log:
        logger.info(f"[PDF移動] {len(move_log)} 件を processed/ へ移動")

    report_name = f"report_{now.strftime('%Y%m%d')}.txt"
    report_path = os.path.join(excel_dir, report_name)
    generate_report(
        report_path, success_list, no_name_pdfs, no_match_pdfs,
        multi_match_pdfs, unmatched_excels, backup_log, move_log,
        protected_log, dry_run=dry_run,
    )
    logger.info(f"\n[レポート] {report_path}")

    logger.info(f"\n{'='*60}")
    logger.info("処理完了")
    logger.info(f"{'='*60}")


def main():
    parser = argparse.ArgumentParser(
        description="PDF利用者情報 → Excel ヒアリングシート 自動転記"
    )
    parser.add_argument("--config", default="config.json", help="設定ファイルパス")
    parser.add_argument("--pdf-dir", default="input_pdf", help="PDF入力ディレクトリ")
    parser.add_argument("--excel-dir", default="excel_sheets", help="Excelファイルディレクトリ")
    parser.add_argument("--form", default="form1",
                        help="書き込み先フォーム名 (デフォルト: form1)")
    parser.add_argument("--dry-run", action="store_true", help="ドライラン（書き込みせず確認のみ）")
    parser.add_argument("--log-file", default=None, help="ログファイルパス")
    args = parser.parse_args()

    process_all(
        config_path=args.config,
        pdf_dir=args.pdf_dir,
        excel_dir=args.excel_dir,
        form_name=args.form,
        dry_run=args.dry_run,
        log_file=args.log_file,
    )


if __name__ == "__main__":
    main()
